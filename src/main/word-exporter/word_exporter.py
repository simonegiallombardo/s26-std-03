import os
from disassembly_loader import build_guide
from docx import Document 
from docx.shared import Inches

    #parameters:
        # json_path (str): file path to the json input file
        # depth (DepthSpec, see the loader): max depth level of the disassemblying process
        # include_bom (bool): whether the bill of material has to be added (or not) to the generated docx
    # return value:
        # str: absolute path of the generated docx file

def export_to_word(json_path, depth=None, include_bom=False):
    
    # Loading of the guide with the following choice that can be made in the gui (so: "include or don't include bom" and "choose the amount of depth levels you want to disassemble")
    guide = build_guide(json_path, depth=depth, include_bom=include_bom)

    # Word document creation
    doc = Document()
    doc.add_heading(f"Disassembly guide: {guide.product.name}", 0)

    # printing the validation warnings taken from the loader
    if guide.warnings:
        doc.add_heading("Validation warnings", level=1)
        for warning in guide.warnings:
            p_warn = doc.add_paragraph(style='Normal')
            r_warn = p_warn.add_run(f"[{warning.severity.upper()}] Rule '{warning.rule}': {warning.message}")
            r_warn.italic = True
    
    # printing the bill of materials (only if requested and present )
    if hasattr(guide, 'bill_of_materials') and guide.bill_of_materials:
        doc.add_heading("Bill of materials", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        
        # column headers
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Component'
        hdr_cells[1].text = 'Weight'
        hdr_cells[2].text = 'Material'
        hdr_cells[3].text = 'Color'
        
        # make the header text bold
        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        
        # iterate over the bom items
        for item in guide.bill_of_materials:
            row_cells = table.add_row().cells
            row_cells[0].text = str(item.name)
            
            # printing the weight value (and removing the ".0" if it's an integer)
            if item.weight is not None:
                row_cells[1].text = f"{int(item.weight) if isinstance(item.weight, float) and item.weight.is_integer() else item.weight} g"
            else:
                row_cells[1].text = "—"
             
            # insert material and color, if they are present   
            row_cells[2].text = str(item.material) if item.material is not None else "—"
            row_cells[3].text = str(item.color) if item.color is not None else "—"
        
        # set column widths 
        col_widths = [Inches(3.0), Inches(1.0), Inches(1.0), Inches(1.0)]
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width
        doc.add_paragraph()
    
    # loop through all next steps to collect the names of the sub-roots that will be disassembled later (in order to print "Then resume from ..." message)
    active_subroots = set()
    for step in guide.steps:
        source_name = guide.product.name if step.index == 1 or not guide.steps[step.index - 2].continues_as else guide.steps[step.index - 2].continues_as.name
        active_subroots.add(source_name)
    
    for step in guide.steps:
        # current object (root or sub-root) we are disassembling
        current_source = guide.product.name if step.index == 1 or not guide.steps[step.index - 2].continues_as else guide.steps[step.index - 2].continues_as.name
        doc.add_paragraph() 
        doc.add_heading(f"Disassembling from {current_source}", level=1)
        
        # print action
        p_action = doc.add_paragraph()
        r_action = p_action.add_run(f"Action {step.index}: {step.operation}")
        r_action.bold = True
        
        # print required tools (if they exist)
        if step.tools_required:
            p_tools = doc.add_paragraph()
            r_tools_label = p_tools.add_run("Tools required: ")
            r_tools_label.bold = True
            p_tools.add_run(", ".join(step.tools_required))
        
        # print steps as a list
        if step.actions:
            for action in step.actions:
                doc.add_paragraph(action.text, style='List Bullet 2')

        collected_sub_roots = []
        
        # if the composite object that i have obtained can be still disassembled, i add it to the sub-roots list
        if step.continues_as and step.continues_as.name in active_subroots:
            collected_sub_roots.append(step.continues_as.name)
            
        # if, for each of the objects i removed from the composite, those object can be still disassembled, i add them to the sub-roots list
        for out_component in step.outputs:
            if out_component.name in active_subroots and out_component.name not in collected_sub_roots:
                collected_sub_roots.append(out_component.name)
        
        # printing "Then resume from <dismountable sub-root component>"
        if collected_sub_roots:
            p = doc.add_paragraph()
            resume_phrases = [f'"Disassembling from {name}"' for name in collected_sub_roots]
            r = p.add_run("Then resume from " + ", ".join(resume_phrases))
            r.bold = True

    # save file to disk
    base_name = os.path.splitext(os.path.basename(json_path))[0]
    output_path = f"{base_name}.docx"
    doc.save(output_path)
    return os.path.abspath(output_path)
